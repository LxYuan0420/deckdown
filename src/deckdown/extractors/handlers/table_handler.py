from __future__ import annotations

from typing import Any, Optional

from pptx.enum.shapes import MSO_SHAPE_TYPE

from pptx.oxml.ns import qn
from deckdown.ast import ShapeKind, TableCell, TablePayload, TableShape, TextPayload
from deckdown.extractors.context import ExtractContext
from deckdown.extractors.handlers.base import ShapeHandler
from deckdown.extractors.utils import extract_text_payload


class TableShapeHandler(ShapeHandler):
    def supports(self, shape: Any) -> bool:  # noqa: ANN401
        return getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.TABLE and bool(
            getattr(shape, "has_table", False)
        )

    def build(self, shape: Any, *, z: int, ctx: ExtractContext) -> Optional[TableShape]:  # noqa: ANN401
        bbox = ctx.bbox(
            left_emu=int(getattr(shape, "left", 0)),
            top_emu=int(getattr(shape, "top", 0)),
            width_emu=int(getattr(shape, "width", 0)),
            height_emu=int(getattr(shape, "height", 0)),
        )
        tbl = shape.table
        n_rows = len(tbl.rows)
        n_cols = len(tbl.columns)
        colspans = [[1 for _ in range(n_cols)] for _ in range(n_rows)]
        rowspans = [[1 for _ in range(n_cols)] for _ in range(n_rows)]
        hmerge = [[False for _ in range(n_cols)] for _ in range(n_rows)]
        vmerge = [[False for _ in range(n_cols)] for _ in range(n_rows)]
        for r in range(n_rows):
            for c in range(n_cols):
                tc = tbl.cell(r, c)._tc
                gs = tc.get("gridSpan")
                rs = tc.get("rowSpan")
                hm = tc.get("hMerge")
                vm = tc.get("vMerge")
                if gs:
                    try:
                        colspans[r][c] = max(1, int(gs))
                    except Exception:
                        pass
                if rs:
                    try:
                        rowspans[r][c] = max(1, int(rs))
                    except Exception:
                        pass
                if hm:
                    hmerge[r][c] = True
                if vm:
                    vmerge[r][c] = True

        visited: set[tuple[int, int]] = set()
        out_cells: list[TableCell] = []
        for r in range(n_rows):
            for c in range(n_cols):
                if (r, c) in visited:
                    continue
                if hmerge[r][c] or vmerge[r][c]:
                    continue
                rowspan = max(1, rowspans[r][c])
                colspan = max(1, colspans[r][c])
                for rr in range(r, min(r + rowspan, n_rows)):
                    for cc in range(c, min(c + colspan, n_cols)):
                        if rr == r and cc == c:
                            continue
                        visited.add((rr, cc))
                text: TextPayload = extract_text_payload(tbl.cell(r, c).text_frame)
                out_cells.append(TableCell(r=r, c=c, rowspan=rowspan, colspan=colspan, text=text))

        payload = TablePayload(rows=n_rows, cols=n_cols, cells=tuple(out_cells))
        return TableShape(
            id=f"s{getattr(shape, 'shape_id', z)}",
            kind=ShapeKind.TABLE,
            name=getattr(shape, "name", None),
            bbox=bbox,
            z=z,
            rotation=None,
            table=payload,
        )

